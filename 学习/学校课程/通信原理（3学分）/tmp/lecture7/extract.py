from pathlib import Path
from pptx import Presentation
from docx import Document
root=Path.cwd()
ppt=next(p for p in (root/'PPT'/'第七讲').glob('*.pptx') if not p.name.startswith('~$'))
prs=Presentation(ppt)
out=[]
for i,s in enumerate(prs.slides,1):
    out.append(f'\n--- PAGE {i} ---\n')
    for sh in s.shapes:
        if sh.has_text_frame: out.append(sh.text)
        if sh.has_table:
            out.extend(' | '.join(c.text for c in row.cells) for row in sh.table.rows)
(root/'tmp'/'lecture7'/'slide-text.txt').write_text('\n'.join(out),encoding='utf-8')
d=Document(next((root/'PPT'/'第七讲').glob('*.docx')))
t='\n'.join(p.text for p in d.paragraphs)
for tb in d.tables:
    t+='\n'+'\n'.join(' | '.join(c.text for c in row.cells) for row in tb.rows)
(root/'tmp'/'lecture7'/'exercise-text.txt').write_text(t,encoding='utf-8')
print('SLIDES:',len(prs.slides))
print('\n'.join(out))
print('\n--- WORKSHEET ---\n'+t)
